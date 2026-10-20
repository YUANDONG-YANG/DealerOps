#!/usr/bin/env python3
"""Generate the DealerOps Group Assignment 2 test case workbook (55 cases)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/Users/robin/IdeaProjects/DealerOps/CPRG-305-B/group-assignment-2/DealerOps_Group_Assignment_2_Test_Cases.xlsx"

HEADERS = [
    "Test Case ID", "Requirement ID", "Test Description", "Pre-Conditions",
    "Test Steps", "Expected Result", "Actual Result", "Test Status",
    "Environment", "Tested By", "Test Date", "Comments", "Priority", "Severity",
]

TESTERS = {
    "QA": "Yuandong Yang",
    "GW": "Logan Jones",
    "AI": "Jackson Warga",
    "UI": "Bedgel Fadhil Ndam Woukouo",
}

ENV_REVIEW = "Dev (static review)"
ENV_UNIT = "Dev (unit, JUnit)"
ENV_INT = "Dev (integration, MockMvc + Testcontainers MySQL)"
ENV_SYS = "Staging (gateway 8080 + core 8081 + ai 8082 + dealer-web)"
ENV_ACC = "Staging (acceptance walkthrough with synthetic dealership data)"

cases = []

# Static review cases (ST-01 to ST-05)
cases += [
    ("ST-01", "REQ-AUTH", "Review authentication and JWT configuration in dealer-core and dealer-gateway",
     "Source for dealer-core SecurityConfig and gateway JwtAuthenticationFilter is available",
     "1. Open dealer-core SecurityConfig and JwtTokenProvider\n2. Open dealer-gateway JwtAuthenticationFilter\n3. Verify shared JWT secret source and expiry policy\n4. Verify all protected routes require authentication",
     "No unprotected business route; JWT secret sourced from environment; expiry enforced consistently",
     "", "Pass/Fail", ENV_REVIEW, TESTERS["QA"], "", "Static review — chapter 5 traceability", "High", "Critical"),
    ("ST-02", "REQ-TENANT", "Review dealer membership and tenant isolation enforcement",
     "DealerMembership entity, service, and controller source available",
     "1. Trace dealerId resolution from JWT claim through service layer\n2. Confirm no controller trusts request-body dealerId\n3. Confirm repository queries filter by authenticated dealerId",
     "All data access scoped by authenticated membership; client-supplied dealerId is ignored or rejected",
     "", "Pass/Fail", ENV_REVIEW, TESTERS["QA"], "", "Static review — tenant boundary", "High", "Critical"),
    ("ST-03", "REQ-INV", "Review inventory lifecycle state machine (DRAFT→READY→SOLD) and sold-lock guard",
     "VehicleService, VehicleStatus enum, and listing controller source available",
     "1. Walk through status transitions in VehicleService\n2. Confirm sold vehicles reject mutating endpoints\n3. Confirm stale-readiness invalidation when vehicle details change",
     "Transitions are explicit; SOLD vehicles rejected on edit; stale readiness recomputed",
     "", "Pass/Fail", ENV_REVIEW, TESTERS["QA"], "", "Static review — inventory", "High", "Major"),
    ("ST-04", "REQ-AI", "Review ai-service internal-only enforcement and prompt boundaries",
     "ai-service source, InternalHeaderFilter, and prompt builder available",
     "1. Verify X-Dealer-Internal header is required on all AI endpoints\n2. Review prompt builder to confirm no secrets or PII templated\n3. Review timeout and retry configuration",
     "AI endpoints reject external calls; prompts contain only structured facts; timeouts bounded",
     "", "Pass/Fail", ENV_REVIEW, TESTERS["AI"], "", "Static review — AI boundary", "High", "Critical"),
    ("ST-05", "REQ-GW", "Review gateway route table, CORS, and timeout configuration",
     "dealer-gateway application.yml and route configuration available",
     "1. Inventory all public routes exposed via 8080\n2. Verify no direct route to ai-service 8082 is published\n3. Verify CORS allowed origins restrict to configured frontend\n4. Verify per-route timeout and circuit breaker",
     "Only intended business routes exposed; AI port not reachable from browser; CORS restrictive; timeouts set",
     "", "Pass/Fail", ENV_REVIEW, TESTERS["GW"], "", "Static review — gateway", "High", "Critical"),
]

# Unit cases (UT-01 to UT-10)
cases += [
    ("UT-01", "REQ-AUTH", "Password hashing uses BCrypt with configured strength",
     "dealer-core unit test environment",
     "1. Call PasswordEncoder.encode(\"P@ssw0rd!\")\n2. Verify output starts with $2a$ or $2b$\n3. Verify encoder.matches raw/encoded",
     "Hash is BCrypt and raw password matches; different calls yield different hashes",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "", "High", "Critical"),
    ("UT-02", "REQ-AUTH", "JwtTokenProvider rejects expired tokens",
     "JwtTokenProvider bean available with test secret",
     "1. Generate token with expiry = now-1s\n2. Call validateToken(token)",
     "validateToken returns false / throws expired token exception",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "", "High", "Critical"),
    ("UT-03", "REQ-TENANT", "DealerContext resolver returns dealerId from authenticated principal only",
     "DealerContext resolver with mocked security context",
     "1. Set SecurityContext with principal containing dealerId=A\n2. Call resolve()\n3. Repeat with no principal",
     "Returns 'A' with principal; throws Unauthorized when principal is absent; ignores request-provided dealerId",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "", "High", "Critical"),
    ("UT-04", "REQ-INV", "Vehicle validator rejects invalid VIN length",
     "VehicleValidator unit scope",
     "1. Validate vehicle with VIN of 16 chars\n2. Validate vehicle with VIN of 17 valid chars",
     "16-char VIN rejected with VIN_INVALID; 17-char valid VIN accepted",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "Boundary check", "High", "Major"),
    ("UT-05", "REQ-INV", "VehicleService.update rejects edits on SOLD vehicle",
     "VehicleService with mocked repository returning SOLD vehicle",
     "1. Call update(id, newPrice) on SOLD vehicle",
     "Throws VehicleLockedException with status=SOLD",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "", "High", "Critical"),
    ("UT-06", "REQ-COMP", "Compliance evaluator flags listing without current evidence as NOT_READY",
     "ComplianceEvaluator unit scope",
     "1. Build listing with no compliance evidence\n2. Call evaluate(listing)",
     "Result = NOT_READY with reason MISSING_COMPLIANCE",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "", "High", "Critical"),
    ("UT-07", "REQ-COMP", "Compliance evidence invalidated when key vehicle fields change",
     "Listing with PASSED evidence; evaluator unit scope",
     "1. Mutate vehicle.year on the linked vehicle\n2. Re-evaluate listing readiness",
     "Evaluator reports STALE_EVIDENCE and listing is not Ready/exportable",
     "", "Pass/Fail", ENV_UNIT, TESTERS["GW"], "", "", "High", "Major"),
    ("UT-08", "REQ-AI", "AI client maps upstream 504 to AiTimeoutException",
     "AI WebClient configured with 2s timeout and mock server returning delay",
     "1. Call AiClient.check(request) against delaying stub\n2. Assert exception type",
     "AiTimeoutException raised; no partial result stored",
     "", "Pass/Fail", ENV_UNIT, TESTERS["AI"], "", "Timeout mapping", "High", "Major"),
    ("UT-09", "REQ-AI", "AI prompt builder excludes customer PII fields",
     "PromptBuilder with sample listing including customer phone and email",
     "1. Build prompt for ad-check on listing with embedded customer contact\n2. Inspect generated prompt text",
     "Prompt contains only vehicle facts; no customer email, phone, or secret present",
     "", "Pass/Fail", ENV_UNIT, TESTERS["AI"], "", "No-leak check", "High", "Critical"),
    ("UT-10", "REQ-AI", "AI client returns BLOCK verdict unchanged when provider output is BLOCK",
     "Mock provider returning BLOCK verdict with reason",
     "1. Call AiClient.check(request)\n2. Inspect returned verdict and reasons",
     "Verdict = BLOCK; reasons preserved verbatim; no silent conversion to PASS",
     "", "Pass/Fail", ENV_UNIT, TESTERS["AI"], "", "", "High", "Critical"),
]

# Integration cases (IT-01 to IT-20)
cases += [
    ("IT-01", "REQ-AUTH", "POST /api/auth/login with valid credentials returns 200 and JWT",
     "Seed user admin@dealer-a.example with known password exists",
     "1. POST /api/auth/login { username, password }",
     "HTTP 200; response contains accessToken (JWT); claims include dealerId",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-02", "REQ-AUTH", "POST /api/auth/login with wrong password returns 401",
     "Seed user exists",
     "1. POST /api/auth/login with incorrect password",
     "HTTP 401; no token; response does not reveal whether username exists",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-03", "REQ-AUTH", "GET /api/vehicles without Authorization header returns 401",
     "Gateway and core running",
     "1. GET /api/vehicles without token",
     "HTTP 401; no vehicle data in body",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-04", "REQ-TENANT", "User from Dealer A cannot read Dealer B vehicle by ID",
     "Dealer A and B seeded; vehicle V-B123 belongs to Dealer B",
     "1. Login as user-A\n2. GET /api/vehicles/V-B123",
     "HTTP 404 or 403; no vehicle details for V-B123 returned",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "Cross-tenant disclosure test", "High", "Critical"),
    ("IT-05", "REQ-TENANT", "Request body dealerId does not switch tenant context",
     "Login as user-A; vehicle V-B123 belongs to Dealer B",
     "1. POST /api/vehicles with body dealerId=B and target V-B123 fields",
     "HTTP 400/403; no record created under Dealer B; audit log records attempt",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-06", "REQ-INV", "POST /api/vehicles creates vehicle with status=DRAFT",
     "Login as admin-A",
     "1. POST /api/vehicles with valid VIN and vehicle fields",
     "HTTP 201; response status=DRAFT; vehicle visible on GET /api/vehicles under Dealer A",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Major"),
    ("IT-07", "REQ-INV", "PUT /api/vehicles/{id} on SOLD vehicle returns 409",
     "Vehicle V-A200 is SOLD",
     "1. PUT /api/vehicles/V-A200 with price change",
     "HTTP 409 with code=VEHICLE_SOLD_LOCKED; price unchanged in DB",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-08", "REQ-INV", "POST /api/vehicles with VIN length 16 returns 400 VIN_INVALID",
     "Login as admin-A",
     "1. POST /api/vehicles with VIN of 16 chars",
     "HTTP 400; error code=VIN_INVALID; no record created",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Major"),
    ("IT-09", "REQ-INV", "POST /api/vehicles/{id}/photos rejects >5MB file",
     "Login as admin-A; vehicle V-A101 in DRAFT",
     "1. POST photo of 6MB JPEG",
     "HTTP 413 or 400; code=PHOTO_TOO_LARGE; no photo persisted",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "Medium", "Major"),
    ("IT-10", "REQ-INV", "POST /api/vehicles/{id}/photos rejects non-image mime type",
     "Login as admin-A; vehicle V-A101 in DRAFT",
     "1. POST a .exe file as multipart image",
     "HTTP 400; code=PHOTO_UNSUPPORTED_TYPE; no photo persisted",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "Medium", "Major"),
    ("IT-11", "REQ-COMP", "POST /api/listings/{id}/ready without compliance evidence returns 409",
     "Listing L-A1 has no compliance evidence",
     "1. POST /api/listings/L-A1/ready",
     "HTTP 409; code=COMPLIANCE_REQUIRED; status unchanged",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-12", "REQ-COMP", "Changing listing price after PASS evidence invalidates readiness",
     "Listing L-A2 has current PASS evidence and status=READY",
     "1. PATCH /api/listings/L-A2 { price: newPrice }\n2. GET /api/listings/L-A2",
     "Status transitions to NOT_READY with reason=STALE_EVIDENCE",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Major"),
    ("IT-13", "REQ-COMP", "GET /api/listings/{id}/export returns 409 when NOT_READY",
     "Listing L-A1 is NOT_READY",
     "1. GET /api/listings/L-A1/export",
     "HTTP 409; code=LISTING_NOT_READY; no export artifact produced",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
    ("IT-14", "REQ-AI", "Ad-check happy path returns PASS verdict with reasons",
     "Listing L-A2 is READY; AI provider stub returns PASS",
     "1. POST /api/listings/L-A2/ad-check",
     "HTTP 200; verdict=PASS; reasons list present; no secret or PII in response",
     "", "Pass/Fail", ENV_INT, TESTERS["AI"], "", "", "High", "Major"),
    ("IT-15", "REQ-AI", "Ad-check with provider timeout maps to HTTP 504",
     "AI stub configured to delay 20s; AI timeout=2s",
     "1. POST /api/listings/L-A2/ad-check",
     "HTTP 504; code=AI_TIMEOUT; no partial verdict stored",
     "", "Pass/Fail", ENV_INT, TESTERS["AI"], "", "Failure-safe behavior", "High", "Major"),
    ("IT-16", "REQ-AI", "Direct call to ai-service port 8082 from outside is blocked",
     "ai-service running; firewall rules applied",
     "1. From external host, curl http://host:8082/api/ai/health",
     "Connection refused or 403; no response body from ai-service",
     "", "Pass/Fail", ENV_INT, TESTERS["AI"], "", "Internal-only enforcement", "High", "Critical"),
    ("IT-17", "REQ-GW", "Gateway health endpoint returns 200 and status=UP",
     "Gateway running",
     "1. GET /actuator/health",
     "HTTP 200; status=UP; subcomponents listed",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "Medium", "Minor"),
    ("IT-18", "REQ-GW", "Gateway rejects request with forged X-Dealer-Internal header",
     "User-A logged in",
     "1. GET /api/vehicles with header X-Dealer-Internal: 1",
     "HTTP 400 or header stripped; request routed normally with user's own dealerId",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "Boundary spoofing attempt", "High", "Critical"),
    ("IT-19", "REQ-CRM", "Customer creation records audit entry with actor and dealerId",
     "Login as admin-A",
     "1. POST /api/customers { name, email }\n2. GET /api/audit?entity=customer",
     "Customer created under Dealer A; audit entry contains actor=admin-A, dealerId=A, action=CREATE",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "Medium", "Major"),
    ("IT-20", "REQ-CRM", "Work order linked to another dealer's customer is rejected",
     "Login as admin-A; customer C-B500 belongs to Dealer B",
     "1. POST /api/work-orders { customerId: 'C-B500' }",
     "HTTP 400/403; code=TENANT_MISMATCH; no work order created",
     "", "Pass/Fail", ENV_INT, TESTERS["GW"], "", "", "High", "Critical"),
]

# System cases (SYS-01 to SYS-10)
cases += [
    ("SYS-01", "REQ-INV", "End-to-end: create vehicle via web, verify in list, edit, sell, confirm locked",
     "dealer-web served through gateway; admin-A logged in",
     "1. Click 'Add vehicle' and submit form\n2. Verify vehicle appears in inventory list\n3. Edit price and save\n4. Mark vehicle SOLD\n5. Attempt second edit",
     "Vehicle created; edits persist; sold transition succeeds; subsequent edit UI shows locked and API rejects",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "High", "Major"),
    ("SYS-02", "REQ-INV", "Upload a photo, verify preview and persistence after refresh",
     "Vehicle V-A101 in DRAFT; user admin-A logged in",
     "1. Open vehicle detail\n2. Upload a 2MB JPEG\n3. Refresh page",
     "Photo preview shows; thumbnail persists after refresh; stored copy retrievable via API",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "Medium", "Major"),
    ("SYS-03", "REQ-COMP", "End-to-end listing workflow: compliance pass → Ready → export",
     "Listing L-A2 with vehicle data populated",
     "1. Trigger compliance check\n2. Verify PASS banner\n3. Click 'Mark Ready'\n4. Click 'Export'",
     "Readiness check PASS; Ready status set; export generates artifact and recorded in listing history",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "High", "Critical"),
    ("SYS-04", "REQ-AI", "Ad-check fail-safe when AI provider is unreachable",
     "ai-service configured with unreachable provider; listing L-A2 Ready",
     "1. Click 'Run ad check' in UI\n2. Wait for completion",
     "UI shows 'AI unavailable' banner; no false PASS stored; export button not enabled by AI result",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "Failure-safe UI latch", "High", "Critical"),
    ("SYS-05", "REQ-GW", "Browser can only reach gateway (8080); direct 8081/8082 blocked",
     "Services running on their own ports",
     "1. Open browser devtools and attempt fetch to http://host:8081/api/vehicles\n2. Attempt fetch to http://host:8082/api/ai/health",
     "Both direct calls fail (CORS or connection); only gateway responds on 8080",
     "", "Pass/Fail", ENV_SYS, TESTERS["GW"], "", "", "High", "Critical"),
    ("SYS-06", "REQ-CRM", "Audit trail: user action logged with user, dealer, timestamp",
     "admin-A logged in",
     "1. Create customer\n2. Open audit log page\n3. Filter by today's date",
     "Audit entry visible with actor=admin-A, dealerId=A, action=CREATE, timestamp within last minute",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "Medium", "Major"),
    ("SYS-07", "REQ-WEB", "Session expiry: expired JWT prompts re-login and preserves destination",
     "admin-A logged in with JWT expiring in 10 seconds",
     "1. Wait 15 seconds\n2. Click inventory menu",
     "Request returns 401; UI redirects to login; after login, returns to inventory page",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "Medium", "Major"),
    ("SYS-08", "REQ-WEB", "UI shows read-only banner when viewing SOLD vehicle",
     "Vehicle V-A200 is SOLD; admin-A logged in",
     "1. Open V-A200 detail page",
     "Edit controls disabled; banner reads 'Vehicle sold — read only'; save button hidden",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "Medium", "Minor"),
    ("SYS-09", "REQ-WEB", "Form validation surfaces API error codes to user clearly",
     "admin-A logged in; VIN length 16",
     "1. Submit vehicle form with 16-char VIN",
     "Form shows inline error 'VIN must be 17 characters'; no vehicle created",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "Medium", "Minor"),
    ("SYS-10", "REQ-WEB", "Logout clears session and token from browser storage",
     "admin-A logged in",
     "1. Click 'Logout'\n2. Inspect browser localStorage / sessionStorage",
     "No access token or dealer context remaining; next API call returns 401",
     "", "Pass/Fail", ENV_SYS, TESTERS["UI"], "", "", "Medium", "Major"),
]

# Acceptance cases (AT-01 to AT-10)
cases += [
    ("AT-01", "REQ-AUTH", "Dealer admin signs in and sees own dealership context",
     "Admin credentials for Dealer A on hand",
     "1. Navigate to login page\n2. Enter admin credentials\n3. Observe landing dashboard",
     "Admin lands on dashboard displaying 'Dealer A' context and admin menu",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "High", "Critical"),
    ("AT-02", "REQ-TENANT", "Dealer A user cannot see Dealer B listings anywhere in UI",
     "Dealer A user logged in; Dealer B has listings",
     "1. Open inventory, listings, customers, audit pages\n2. Search for Dealer B VINs or names",
     "No Dealer B records visible or searchable; API audit shows no cross-tenant read attempt succeeded",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "High", "Critical"),
    ("AT-03", "REQ-COMP", "Listing cannot be exported without current PASSED compliance",
     "Listing L-A3 with stale evidence",
     "1. Open listing L-A3\n2. Attempt to export",
     "Export button disabled or export API returns 409; user sees reason 'Compliance evidence stale'",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "High", "Critical"),
    ("AT-04", "REQ-AI", "Ad check never reports PASS when provider is unavailable",
     "AI provider offline",
     "1. Run ad check on listing L-A2",
     "Verdict shown as 'Unavailable, try later'; no PASS recorded in listing history",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "High", "Critical"),
    ("AT-05", "REQ-CRM", "Audit log can be filtered by user and action for compliance review",
     "At least 10 audit records exist for Dealer A",
     "1. Open audit page\n2. Filter by actor=admin-A and action=UPDATE",
     "Filtered list shows only matching records; export CSV contains filtered rows only",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "Medium", "Major"),
    ("AT-06", "REQ-WEB", "Full lifecycle: create vehicle, add photo, compliance check, mark ready, export",
     "admin-A logged in; fresh VIN available",
     "1. Create vehicle → upload photo → set price → run compliance → mark ready → export",
     "All steps succeed; export artifact lists correct vehicle data and compliance evidence ID",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "Golden path", "High", "Critical"),
    ("AT-07", "REQ-WEB", "Non-admin user cannot access admin-only menus",
     "Regular sales user for Dealer A logged in",
     "1. Attempt to open 'Users' and 'Audit' menus",
     "Menus hidden; direct URL access returns 403; no admin data leaked",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "High", "Critical"),
    ("AT-08", "REQ-WEB", "Error messages are user-readable and do not leak stack traces",
     "admin-A logged in",
     "1. Trigger validation failure (bad VIN)\n2. Trigger 409 (sold vehicle edit)",
     "Both show business-friendly messages; no Java stack trace, no SQL, no secret in UI or network body",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "Medium", "Major"),
    ("AT-09", "REQ-WEB", "Browser compatibility: workflow succeeds in Chrome and one other Chromium browser",
     "Two supported browsers installed",
     "1. Repeat AT-06 in Chrome\n2. Repeat in secondary browser",
     "Both complete without layout break or console error",
     "", "Pass/Fail", ENV_ACC, TESTERS["UI"], "", "", "Medium", "Minor"),
    ("AT-10", "REQ-WEB", "After acceptance run, no real customer PII exists in test data or evidence",
     "Evidence folder for this acceptance run available",
     "1. Grep evidence for real email domains outside test allowlist\n2. Review screenshots for PII",
     "No real PII detected; test data uses synthetic names/emails/VINs only",
     "", "Pass/Fail", ENV_ACC, TESTERS["QA"], "", "Pre-submission scrub", "High", "Critical"),
]

assert len(cases) == 55, f"Expected 55 cases, got {len(cases)}"

wb = Workbook()
ws = wb.active
ws.title = "Test Cases"

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F4E78")
thin = Side(border_style="thin", color="B7C9D6")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")

for i, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = border

for r, row in enumerate(cases, start=2):
    for c_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c_idx, value=val)
        cell.alignment = wrap
        cell.border = border

widths = [12, 14, 36, 32, 44, 38, 20, 12, 28, 22, 14, 24, 10, 12]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 32
ws.freeze_panes = "A2"

# Summary sheet
summary = wb.create_sheet("Summary")
summary["A1"] = "DealerOps Group Assignment 2 — Test Cases"
summary["A1"].font = Font(bold=True, size=14, color="1F4E78")
summary["A3"] = "Group:"
summary["B3"] = "Group 12 (team G2)"
summary["A4"] = "Members:"
summary["B4"] = "Yuandong Yang (Robin), Bedgel Fadhil Ndam Woukouo, Jackson Warga, Logan Jones"
summary["A5"] = "Course:"
summary["B5"] = "CPRG-305-B Software Testing and Deployment (Fall 2026)"
summary["A6"] = "Document version:"
summary["B6"] = "v1.0 (2026-10-13)"
summary["A8"] = "Phase"
summary["B8"] = "ID range"
summary["C8"] = "Count"
for cell in (summary["A8"], summary["B8"], summary["C8"]):
    cell.font = header_font
    cell.fill = header_fill
rows = [
    ("Requirements/design review", "ST-01 – ST-05", 5),
    ("Unit", "UT-01 – UT-10", 10),
    ("Integration", "IT-01 – IT-20", 20),
    ("System", "SYS-01 – SYS-10", 10),
    ("Acceptance", "AT-01 – AT-10", 10),
    ("Total", "", 55),
]
for i, r in enumerate(rows, start=9):
    for j, v in enumerate(r, start=1):
        summary.cell(row=i, column=j, value=v)
summary.column_dimensions["A"].width = 28
summary.column_dimensions["B"].width = 20
summary.column_dimensions["C"].width = 10

wb.save(OUT)
print(OUT)
